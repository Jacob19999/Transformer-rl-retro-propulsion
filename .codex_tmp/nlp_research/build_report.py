from pathlib import Path
from copy import deepcopy
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(r'C:\Transformer-rl-retro-propulsion')
OUT = ROOT / 'Plan' / 'NLP'
OUT.mkdir(parents=True, exist_ok=True)
doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = sec.bottom_margin = Inches(.72)
sec.left_margin = sec.right_margin = Inches(.8)
sec.footer_distance = Inches(.3)
normal = doc.styles['Normal']
normal.font.name = 'Calibri'
normal.font.size = Pt(10.5)
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.line_spacing = 1.07
for sn, size in [('Title', 26), ('Subtitle', 12), ('Heading 1', 18), ('Heading 2', 12)]:
    s = doc.styles[sn]
    s.font.name = 'Calibri'
    s.font.size = Pt(size)
    s.font.color.rgb = RGBColor(0,0,0)
    s.paragraph_format.space_after = Pt(8)
    s.paragraph_format.space_before = Pt(10 if sn != 'Title' else 0)
    s.paragraph_format.keep_with_next = True
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = footer.add_run('NLP project ideas  |  ')
r.font.size = Pt(9)
fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE'); footer._p.append(fld)
doc.core_properties.title = 'Embedding projects for aviation and landing trajectories'
doc.core_properties.subject = 'Research ideas and project scope'
doc.core_properties.author = 'Research planning'

def p(text, boldlead=None):
    a = doc.add_paragraph()
    if boldlead and text.startswith(boldlead):
        a.add_run(boldlead).bold = True; a.add_run(text[len(boldlead):])
    else: a.add_run(text)
    return a

def h(text, level=2): return doc.add_heading(text, level)
def page(title): doc.add_page_break(); h(title,1)
def bullet(text): return doc.add_paragraph(text, 'List Bullet')
def link(par, label, url):
    e=OxmlElement('w:hyperlink'); e.set(qn('r:id'), par.part.relate_to(url, RT.HYPERLINK, is_external=True))
    r=OxmlElement('w:r'); pr=OxmlElement('w:rPr'); c=OxmlElement('w:color'); c.set(qn('w:val'),'24577A'); pr.append(c); r.append(pr)
    t=OxmlElement('w:t'); t.text=label; r.append(t); e.append(r); par._p.append(e)

def table(headers, rows, widths):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    for c,w in zip(t.columns,widths): c.width=Inches(w)
    for c,s in zip(t.rows[0].cells,headers): c.text=s
    rep=OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(rep)
    for row in rows:
        for c,s in zip(t.add_row().cells,row): c.text=str(s)
    for i,row in enumerate(t.rows):
        cant=OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(cant)
        for j,c in enumerate(row.cells):
            c.width=Inches(widths[j]); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp=c._tc.get_or_add_tcPr(); shading=OxmlElement('w:shd'); shading.set(qn('w:fill'),'24475D' if i==0 else ('F0F4F6' if i%2==0 else 'FFFFFF')); cp.append(shading)
            borders=OxmlElement('w:tcBorders')
            for side in ('top','bottom','left','right'):
                b=OxmlElement('w:'+side); b.set(qn('w:val'),'single'); b.set(qn('w:sz'),'4'); b.set(qn('w:color'),'D9D9D9'); borders.append(b)
            cp.append(borders); mar=OxmlElement('w:tcMar')
            for side in ('top','bottom','left','right'):
                m=OxmlElement('w:'+side); m.set(qn('w:w'),'90'); m.set(qn('w:type'),'dxa'); mar.append(m)
            cp.append(mar)
            for a in c.paragraphs:
                a.paragraph_format.space_after=Pt(2); a.paragraph_format.line_spacing=1.0
                for r in a.runs:
                    r.font.size=Pt(9.5)
                    if i==0: r.bold=True; r.font.color.rgb=RGBColor(255,255,255)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)
    return t

# Native Word mathematics. Every fraction, index, and sum is an OMML structure.
def mr(s):
    e=OxmlElement('m:r'); t=OxmlElement('m:t'); t.text=s; e.append(t); return e
def seq(*items):
    out=[]
    for x in items:
        if isinstance(x,list): out.extend(x)
        else: out.append(mr(x) if isinstance(x,str) else x)
    return out
def wrap(name, items):
    e=OxmlElement('m:'+name)
    for x in seq(items): e.append(x)
    return e
def sub(a,b):
    e=OxmlElement('m:sSub'); e.append(wrap('e',a)); e.append(wrap('sub',b)); return e
def sup(a,b):
    e=OxmlElement('m:sSup'); e.append(wrap('e',a)); e.append(wrap('sup',b)); return e
def frac(a,b):
    e=OxmlElement('m:f'); e.append(wrap('num',a)); e.append(wrap('den',b)); return e
def sm(low,high,body):
    e=OxmlElement('m:nary'); pr=OxmlElement('m:naryPr'); ch=OxmlElement('m:chr'); ch.set(qn('m:val'),'∑'); pr.append(ch)
    loc=OxmlElement('m:limLoc'); loc.set(qn('m:val'),'undOvr'); pr.append(loc); e.append(pr)
    e.append(wrap('sub',low)); e.append(wrap('sup',high)); e.append(wrap('e',body)); return e
def eq(*items):
    a=doc.add_paragraph(); a.alignment=WD_ALIGN_PARAGRAPH.CENTER
    a.paragraph_format.space_before=Pt(5); a.paragraph_format.space_after=Pt(9)
    om=OxmlElement('m:oMathPara'); math=OxmlElement('m:oMath')
    for x in seq(*items): math.append(x)
    om.append(math); a._p.append(om)

doc.add_heading('Embedding projects for aviation and landing trajectories',0)
p('Research directions and a practical project shortlist',None).style='Subtitle'
p('16 September 2026')
p('Best fit for this repository: natural-language retrieval of landing trajectories. Best fit for a conventional NLP course: retrieval of similar NASA aviation safety reports. Both can produce a complete, measurable project with frozen text encoders; model training can remain an extension.', 'Best fit for this repository:')
p('The central question is whether an embedding preserves the distinctions that matter in aviation: event order, direction, altitude, landing behavior, and negation. A useful project should test those distinctions against simple baselines, rather than rely on an attractive embedding plot.')
h('Six project directions')
table(['Idea','Core deliverable','Training in core','Relative effort'],[
('1 Landing trajectory search','Query runs in ordinary language','None','Medium'),
('2 Aviation safety analog search','Retrieve relevant ASRS narratives','None','Low to medium'),
('3 Aviation meaning sensitivity','Benchmark numbers and opposites','None','Low'),
('4 Aircraft maneuver search','Search short ADS-B segments','None','Medium to high'),
('5 A vocabulary of landing motion','Compare event representations','None*','Medium'),
('6 Align telemetry with language','Learn a trajectory-to-text space','Required','High')],[1.68,2.8,1.1,1.32])
p('*A manual event vocabulary needs no encoder training. Fitting a clustering codebook or learning Word2Vec/VQ tokens is optional training and must be reported as such.')
h('What counts as an embeddings project')
p('Core work: prepare a corpus, encode it with pretrained models, retrieve or compare examples, and evaluate semantic behavior. TF-IDF/BM25 baselines and optional unsupervised clustering are allowed, but distinguish fitting these from updating a neural encoder. Generation, masked-token reconstruction, forecasting, and PPO retraining are outside the minimum project.')
p('Recommendation: develop Idea 1 and borrow the hard contrast examples from Idea 3. If trajectory collection becomes the bottleneck, choose Idea 2. The strongest contribution is a careful aviation-specific evaluation, not a claim that using a Transformer for trajectories is new. [1, 6, 11]')

page('What the repository already provides')
p('The code and saved traces support an offline representation study now. The following findings refer to the checked working tree on 16 September 2026; research goals in the README are not treated as completed capabilities.')
table(['Asset','Verified observation','Project implication'],[
('Observation contract [L1]','24 base channels; optional wind adds 3','Decode named channels and units before describing motion'),
('Evaluation traces [L2]','Time, episode, action, position, world velocity, before/after observation, contact and done','Reuse existing JSONL for a small prototype'),
('Diagnostic inventory [L3]','5 trace files, 25 episodes, 1,243 sampled rows','Pilot corpus only; sampled rows are not independent episodes'),
('Episode exporter [L4]','CSV steps and JSON metadata with seed and configuration/git hashes','Keep provenance attached to every representation'),
('Controller status [L5]','Feed-forward PPO model exists; GTrXL trainer explicitly lacks a sequence-aware optimizer','Do not require a pretrained GTrXL embedding'),
('Landing criteria [L6]','Contact LANDED plus pad-distance tolerance defines landing success','A LANDED label alone is not a complete success label')],[1.35,2.65,2.9])
h('Important details for trajectory encoding')
p('The base observation includes position error, a wxyz quaternion, body-FRD linear and angular velocity, height, fin angles and rates, normalized motor speed, and contact state. World velocity in the trace is a separate field: do not read body-frame vertical velocity as altitude rate. The optional wind channels are an estimate, not proof of disturbance cause. [L1, L2]')
p('The evaluator samples at a configurable trace interval and also logs terminal steps using pre-reset state. Preserve recorded timestamps and episode boundaries. Do not reconstruct missing high-frequency events from sparse samples or interpolate across resets. Verify action scaling against the exporting path before calling the fifth action “physical throttle.” [L2]')
p('Concrete example: episode 1 of full28m_seed456 is marked LANDED, with touchdown speed 0.1564 m/s and pad distance 0.9529 m. This distinguishes a gentle contact from a precise landing; the success helper has a default pad tolerance of 0.5 m, although the relevant run configuration must determine the actual threshold. [L3, L6]')
h('A bounded collection target')
p('Use the 25 episodes for a schema and captioning pilot. Aim for 300–1,000 independently identified episodes covering checkpoints, seeds, outcomes and documented conditions before a larger benchmark. These are proposed targets, not an existing dataset. Keep the final test on a fixed task; analyze curriculum stages and changed physics configurations as separate strata. The NLP work need not change rewards or controllers.')

page('Idea 1 Search landing trajectories with language')
p('Research question: do physics-derived descriptions plus frozen sentence embeddings retrieve landing behavior more reliably than numeric strings or keyword search?', 'Research question:')
p('A user asks “find a run that drifted away from the pad before correcting late.” The system returns the relevant episodes and time windows, with the original trajectory plots and numeric evidence. This is semantic retrieval over a text representation of telemetry; it does not establish that a text encoder directly understands raw physics.')
h('Minimum project')
p('For each episode, calculate a small set of auditable features: height trend, radial pad error, downward speed, tilt, angular-rate oscillation, and action saturation where its scale is verified. Describe ordered segments using fixed, documented rules, for example “early descent with increasing pad error; later correction with decreasing pad error.” Keep duration and magnitudes. Store full-episode and segment descriptions separately.')
eq('τ = (',sub('x','1'),', …, ',sub('x','T'),'),     d = S(τ),     z = ',frac('f(d)',seq(sub('‖f(d)‖','2'))))
p('Here τ is a sampled trajectory, x is a state/action record, S is the deterministic description function, f is a frozen text encoder, and z is a unit-length embedding. A description is lossy, so retain its source measurements and timestamps.')
eq('s(q, τ) = ',sup('u','⊤'),'z,     u = ',frac('f(q)',sub('‖f(q)‖','2')))
p('For an asymmetric retrieval model, apply its prescribed query/document prefixes before encoding. Rank episodes by their full description, or return the best matching segment with its time span. Keep exact constraints such as “touchdown speed below 0.5 m/s” as numeric filters and evaluate this hybrid separately. [1–3]')
h('Experiments that make it research')
bullet('Compare BM25 on descriptions, frozen MiniLM and E5 on the same descriptions, frozen embeddings of formatted numeric sequences, and standardized physical features for trajectory-to-trajectory retrieval. DTW is a numeric query-by-example baseline, not a direct natural-language retriever.')
bullet('Use 40–60 manually authored queries spanning drift, reversal, descent, tilt and stabilization. Judge relevance from plots and measurements, with queries written independently of caption templates. Include same-outcome/different-behavior distractors.')
bullet('Ablate event order, duration, absolute numbers, qualitative bins and terminal labels. Query paraphrases test language generalization; held-out seeds/checkpoints test trajectory generalization.')
p('Measures: nDCG@10, Recall@10 on a fully judged subset, temporal localization for segment queries, and a categorized error analysis. A negative result is useful if embeddings fail on order or magnitudes. Optional extension: tune a small text adapter on human-checked pairs, or pursue Idea 6. Estimated core effort: 3–4 weeks after the pilot.')

page('Idea 2 Retrieve similar aviation safety reports')
p('Research question: can frozen embeddings retrieve operationally similar incidents when pilots describe the same issue with different vocabulary?', 'Research question:')
p('Example query: “The crew followed an incorrect altitude because two clearances sounded similar.” Return relevant reports with highlighted evidence passages. Focus on retrieval and comparison; a generated safety answer is unnecessary.')
h('Data and scope')
p('NASA ASRS publishes de-identified narratives and analyst-coded fields. The database supports CSV exports, limited to 10,000 incident records per download. Start with 1,000–3,000 records on altitude deviations, communication or approach events. For a small pilot, curated report sets contain 50 records each. Deduplicate by accession identifier and keep related reports together. [4, 5]')
p('Embed narrative passages without appending the analyst’s event category. Reserve those categories for weak evaluation or stratification. Chunk long reports by paragraph, retain report IDs, and compare whole-report pooling with passage retrieval. Do not mistake similarity in airports or aircraft names for similarity in the incident mechanism.')
h('Representation and baseline')
eq('s(q, D) = ',sub('max','c ∈ C(D)'), ' ',sup('u','⊤'),sub('z','c'))
p('C(D) is the set of chunks of report D, u is the normalized query vector, and z with index c is the chunk embedding. Maximum chunk similarity is a simple baseline; test whether it over-favors long reports. Compare with BM25 and a fixed reciprocal-rank fusion of lexical and semantic results.')
h('Evaluation')
p('Create 50–80 queries with 0/1/2 relevance grades: irrelevant, partially related, and closely matching operational issue. Judge pooled top results from every system, plus a random sample, without showing the system identity. Include difficult contrasts such as “cleared altitude was misunderstood” versus “cleared altitude was understood but not maintained.”')
p('Use nDCG@10 and judged Precision@10 as primary metrics. Use Recall@k only when the relevant set is complete enough to justify the denominator. Remove near-duplicates and split by incident group, preferably also date. Analyst tags are useful weak labels, not gold similarity judgments.')
h('Why choose this project')
p('It is the clearest natural-language project, has a public corpus, and can succeed without GPUs or model training. Sentence-BERT established the use of independently computed sentence vectors for efficient similarity comparison. The specific contribution here would be an aviation relevance benchmark and analysis of abbreviations, negation and incident mechanisms. [1]')
p('ASRS describes reports as unverified accounts; retrieval relevance does not establish causality or estimate aviation incident rates. Optional extension: domain contrastive fine-tuning using adjudicated positive pairs and hard negatives. Estimated core effort: 2–3 weeks. [4]')

page('Idea 3 Test whether aviation embeddings preserve meaning')
p('Research question: can pretrained embeddings distinguish near-identical aviation statements that imply different actions or events?', 'Research question:')
p('This is the most compact project. It can also be an evaluation component for Ideas 1, 2 or 4. The output is a contrastive benchmark and model comparison, not a flight-control system.')
table(['Anchor','Meaning-preserving positive','Hard negative'],[
('Descend to six thousand feet','Reduce altitude to 6000 ft','Climb to six thousand feet'),
('Turn left heading two four zero','Make a left turn to heading 240','Turn right heading two four zero'),
('The descent stopped before the turn','The aircraft leveled, then turned','The aircraft turned, then leveled'),
('The crew did not descend','The crew maintained altitude','The crew descended')],[2.2,2.45,2.25])
p('These are researcher-authored examples, not quotations from ATC recordings. Review each pair for context: for example, “descend to” and “descend by” are not interchangeable. Separate command semantics, observed events and intentions.')
eq('Δ(q, ',sup('d','+'),', ',sup('d','−'),') = s(q, ',sup('d','+'),') − s(q, ',sup('d','−'),')')
eq('A = ',frac('1','N'),sm('i = 1','N',seq('𝟙[',sub('Δ','i'),' > 0]')))
p('Δ is the similarity margin between a correct paraphrase and a meaning-changing negative. A is triplet accuracy across N examples, with ties counted as failures. Report separate results for direction, altitude, units, negation, callsign identity and temporal order. Avoid one aggregate score hiding a systematic failure.')
h('Data and procedure')
p('Build 200–400 reviewed triplets using 40–60 distinct semantic families. Split by family so a sentence with a swapped number cannot leak across partitions. Use original and paraphrased phrasing, digit/spoken-number variants, and unseen values. Compare TF-IDF, frozen MiniLM, frozen E5, and a transparent slot-based comparator for structured commands.')
p('ATCO2 offers manually transcribed ATC speech and entity annotations; its free one-hour research subset is distinct from the larger licensed releases. It can supply an authentic-language evaluation slice after checking the subset’s stated usage terms. Keep the main benchmark independently authored so no speech recognition or paid corpus is required. [7]')
h('Contribution and extension')
p('Test the hypothesis that embeddings capture broad topical similarity more easily than operational equivalence. Report counterexamples, margin distributions and paired confidence intervals. A small adapter with hard negatives is an optional extension; keep held-out phrase families untouched. Estimated core effort: 1–2 weeks for a narrow benchmark, with annotation quality setting the pace.')

page('Ideas 4 and 5 Represent motion as searchable events')
h('Idea 4 Search aircraft maneuver segments')
p('Research question: can natural-language maneuver descriptions retrieve ADS-B segments across different flights? Example: “a descending left turn followed by level flight.” Use one airport or one flight phase, hundreds of short segments, and an explicitly bounded time period. OpenSky offers public scientific datasets; full historical Trino access has institutional eligibility requirements. Obtain a small usable sample before committing. [8]')
p('Convert positions into an appropriate local metric frame; preserve altitude type, timestamp spacing, track angle, speed and vertical-rate units. Use unwrapped track changes to describe ground-track turns. Resample only across acceptable gaps, retain missingness, and exclude noisy segments from gold labels. Motion in ADS-B does not by itself identify a clearance, cause or pilot intent.')
p('Compare text descriptions plus frozen encoders with BM25 and numeric query-by-example retrieval. Evaluate semantic retrieval on manually reviewed segments with entire flights and days held out. Report Recall@k on a fully judged small subset, nDCG@10 and localization overlap. Optional extension: use validated procedure context for US flights, with a historical CIFP cycle matched to the data. FAA data are not a universal procedure map. [8, 9]')
p('Risk and effort: missing data, coordinate handling and acquisition may dominate the NLP. Allow 3–5 weeks after confirming access. Do not begin with global coverage, weather fusion, automatic explanations or full route prediction.')
h('Idea 5 Compare a vocabulary of landing motion')
p('Research question: which representation preserves similarity between landing behaviors: continuous features, named maneuver events, or numeric text? Use the repo corpus and compare trajectory-to-trajectory retrieval, plus a language-query track using the same event descriptions as Idea 1.')
eq(sub('w','t'),' = Q(',sub('x','t−L+1'),', …, ',sub('x','t'),'),     W = (',sub('w','1'),', …, ',sub('w','M'),')')
p('Q maps a time window of L samples into a documented event such as “descent slowing” or “pad error increasing”; W is a sequence of M events. Use time-based windows when sampling intervals differ. Define thresholds from physical scales or the development partition. Keep durations and allow simultaneous events rather than forcing mutually exclusive flight phases.')
p('Compare event unigram/bigram counts, BM25 over ordinary-language event sequences, and frozen sentence embeddings of ordered descriptions. Test sensitivity to time reversal, removed duration, and different bin widths. Assess agreement with human behavior labels and nearest-neighbor retrieval; use 2D projections only for illustration.')
p('Training extension: fit k-means motion words or Word2Vec on a sufficiently large training corpus, then test stability across seeds and quantify information lost by quantization. A VQ-VAE is a larger extension. Arbitrary tokens such as Z17 or <FIN_SAT> have no guaranteed meaning to a frozen language encoder. Estimated core effort: 2–4 weeks; this is closest to representation learning unless the language-query evaluation is retained.')

page('Idea 6 Learn alignment between telemetry and language')
p('Research question: can a small trajectory encoder retrieve runs from text without relying on a hand-written description function at inference time?', 'Research question:')
p('This is the strongest training extension to Idea 1. Pair each trajectory or segment with descriptions, freeze the text encoder, and learn a compact temporal encoder with a projection into the text embedding space. Start only after a reliable frozen-embedding baseline and independent relevance set exist.')
eq(sub('z','i'),' = ',frac(seq(sub('g','θ'),'(',sub('τ','i'),')'),seq('‖',sub('g','θ'),'(',sub('τ','i'),')',sub('‖','2'))),',     ',sub('u','i'),' = ',frac(seq('f(',sub('d','i'),')'),seq('‖f(',sub('d','i'),')',sub('‖','2'))))
eq('ℒ = −',frac('1','B'),sm('i = 1','B',seq('log ',frac(seq('exp(',sup(sub('z','i'),'⊤'),sub('u','i'),'/η)'),sm('j = 1','B',seq('exp(',sup(sub('z','i'),'⊤'),sub('u','j'),'/η)'))))))
p('g with parameters θ encodes telemetry; f encodes text; B is batch size; η > 0 is temperature. The contrastive objective raises similarity for matched trajectory–description pairs relative to other batch descriptions. This simple one-positive formula assumes the other descriptions are negatives; mask equivalent descriptions or use a multi-positive loss when that assumption fails.')
h('What to train and what to hold fixed')
p('Begin with a small 1D CNN or temporal encoder plus a projection layer. Normalize channels using training statistics, retain timestamps or time deltas, mask padding, and keep physical units in the schema. Reset any recurrent memory at episode boundaries. Do not retrain PPO as part of this extension.')
p('Use independent human descriptions for evaluation. Template-generated pairs may be useful supervision, but a model trained and tested on the same templates may only imitate the event extractor. Contrast same-checkpoint/different-behavior episodes, and different-checkpoint/same-behavior episodes, to expose controller-identity shortcuts.')
h('Evidence and novelty boundary')
p('TMR demonstrated learned text-to-human-motion retrieval with contrastive alignment and a motion-generation objective; it is related methodology, not an aviation model. A June 2026 preprint, CADE, also investigates alignment between time-series representations and frozen text anchors for question answering. These precedents support feasibility of the framing, not success on this repository. [10, 11]')
p('Evaluate both text-to-trajectory and trajectory-to-text retrieval, language paraphrases, unseen run families and a captions-only baseline. A proposed data target is 1,000–5,000 diverse paired segments, with dependence controlled at episode level; actual adequacy must be judged from learning curves. Allow an additional 2–4 weeks and access to a GPU. This is a stretch project, not required for completion.')

page('A defensible experiment and a four week plan')
h('Shared evaluation protocol')
p('Split first, then derive windows and fit preprocessing. Group all windows, captions and paraphrases of one episode or incident together. For repo data, use run/checkpoint and seed families as the strongest available grouping; report when too few independent groups limit generalization. Tune bins, preprocessing and fusion weights on development data only.')
eq('Recall@k = ',frac(seq('|',sub('R','q'),' ∩ ',sub('Top','k'),'(q)|'),seq('|',sub('R','q'),'|')))
eq('nDCG@k = ',frac('1', 'IDCG@k'),sm('r = 1','k',frac(seq(sup('2',sub('rel','r')),' − 1'),seq(sub('log','2'),'(r + 1)'))))
p('R with index q is the complete relevant set for query q. rel with index r is the graded relevance of the result at rank r; IDCG is the ideal ranking’s score. Compute query-wise values and average across queries. Exclude queries with no relevant items from these metrics and report them separately. With incomplete judgments, emphasize judged precision and nDCG under a stated pooling protocol rather than claiming exact recall.')
p('Use paired bootstrap intervals over independent query families, and discuss episode/run dependence for trajectory experiments. Relevance annotation should include a second reviewer on a subset, with disagreements adjudicated. A result is convincing when improvements survive paraphrases and hard negatives, not merely a train/test split of neighboring windows.')
h('Small model and compute choices')
p('Compare sentence-transformers/all-MiniLM-L6-v2 with intfloat/e5-small-v2. Both produce 384-dimensional embeddings. MiniLM’s card describes default truncation beyond 256 wordpieces; E5 supports up to 512 tokens and prescribes query/passage prefixes for retrieval. Record tokenizer and model revision, pooling and normalization. Chunk before truncation and keep the same retrieval units across baselines. [2, 3]')
p('For 10,000 vectors with 384 float32 entries, raw embedding storage is 10,000 × 384 × 4 = 15.36 MB, excluding model weights, text and index overhead. Exact cosine search is adequate for a pilot. A CPU can run these small encoders; report measured throughput instead of assuming a GPU is necessary.')
table(['Week','Idea 1 deliverable','Decision gate'],[
('1','Pilot corpus, provenance, event rules, 15 queries','Can the descriptions be verified from the traces?'),
('2','Larger corpus, BM25, two frozen encoders, query set','Do not expand scope before all baselines run'),
('3','Blind relevance labels, grouped evaluation, ablations','Check leakage and order/number failure modes'),
('4','Error analysis, demo and written results','Add training only if the core study is complete')],[.55,3.35,3.0])
p('Minimum success is a reproducible comparison and an informative answer to the research question. Beating every baseline is not required. A demo should return trace IDs, relevant time windows and source evidence; its purpose is to make the evaluation inspectable.')

page('How to narrow the earlier proposal')
h('Keep the representation comparison')
p('Comparing grid/numeric representations, named events and learned motion words is a sensible research axis. For this scope, make retrieval quality and semantic fidelity the targets. ADE/FDE and altitude prediction error belong to a forecasting extension; they do not directly measure whether a frozen embedding is useful.')
h('Avoid assuming a pretrained aviation language')
p('A Transformer architecture does not automatically make a task NLP, and arbitrary waypoint IDs or numeric coordinate strings do not inherit useful semantics from a text model. Use actual language queries, genuine narratives, or explicit text–trajectory alignment to make the language contribution clear. “FlightBERT” and “FlightBERT++” already identify flight-trajectory prediction research; use a descriptive project title rather than implying a new model family. [6]')
h('Use a score that matches the model')
p('A sentence embedding provides a vector, not a normalized next-token probability. Perplexity therefore requires a probabilistic sequence model and cannot be inferred directly from cosine similarity. An embedding anomaly score can instead be the mean cosine distance to reference neighbors, with calibration on held-out ordinary behavior:')
eq('a(z) = ',frac('1','k'),sm('j ∈ Nₖ(z)','',seq('(1 − ',sup('z','⊤'),sub('z','j'),')')))
p('Exclude self-matches and same-episode near-duplicates. Call this a novelty or outlier score. It does not imply unsafe behavior: a go-around, holding pattern or unusual landing correction may be entirely appropriate. A BERT-style masked model’s pseudo-likelihood is also distinct from autoregressive perplexity.')
h('Treat navigation and multimodal data as later work')
p('CIFP contains navigation and procedure records on a 28-day cycle, but nearest-waypoint matching does not establish the procedure flown or the clearance issued. A valid map matcher must account for direction, geometry, sequence and data vintage. Start with observed kinematics. [9]')
p('Do not add independently trained weather, text and trajectory vectors merely because their dimensions match. Their coordinates need a shared alignment; otherwise use separately evaluated retrieval scores or a learned fusion model. Associating weather with a deviation is also not a causal explanation.')
p('ATC speech plus synchronized surveillance, VQ codebooks, trajectory generation and weather/NOTAM explanation each introduce substantial data or training work. Keep them outside the core deliverable. Aircraft and EDF landing motion can share analysis methods, but their physical scales and operational semantics require separate evaluation.')
h('Suggested project title')
p('Evaluating Frozen Language Embeddings for Semantic Retrieval of Landing Trajectories')
p('Alternative for a purely textual project: Evaluating Semantic Retrieval of Aviation Safety Narratives. For either choice, state the contribution as a scoped benchmark, representation comparison and error analysis; a literature search does not establish that the idea is unprecedented.')

page('Research sources')
p('Primary papers, author model cards and official dataset documentation. Links were checked on 16 September 2026. Bracketed numbers in the report refer to the sources below.')
refs=[
('1','Reimers, N. and Gurevych, I. (2019). Sentence-BERT: Sentence Embeddings using Siamese BERT-Networks. EMNLP-IJCNLP.','https://aclanthology.org/D19-1410/','Foundation for sentence embeddings compared by cosine similarity.'),
('2','Sentence Transformers. all-MiniLM-L6-v2 model card.','https://huggingface.co/sentence-transformers/all-MiniLM-L6-v2','Frozen baseline; model dimensions, pooling example and default input limit.'),
('3','Wang et al. / intfloat. e5-small-v2 model card.','https://huggingface.co/intfloat/e5-small-v2','Frozen retrieval baseline; prefixes, 384 dimensions and token limit.'),
('4','NASA ASRS. ASRS Database Online.','https://asrs.arc.nasa.gov/search/database.html','Narratives, analyst fields, export formats and 10,000-record export limit; reports are not verified.'),
('5','NASA ASRS. Database Report Sets.','https://asrs.arc.nasa.gov/search/reportsets.html','Official topic-based starter corpora; 50 records per report set.'),
('6','Guo, D. et al. (2023; revised 2024). A Non-autoregressive Multi-Horizon Flight Trajectory Prediction Framework with Gray Code Representation.','https://arxiv.org/abs/2305.01658','FlightBERT++ and its relationship to the earlier FlightBERT binary-encoding framework. This is prediction work, not an off-the-shelf text retriever.'),
('7','ATCO2 project. Data available from the project; Zuluaga-Gomez et al. (2022; revised 2023), corpus paper.','https://www.atco2.org/data','Distinguishes the free one-hour research subset, manually annotated four-hour test set and larger licensed corpus.'),
('8','OpenSky Network. Data access and datasets.','https://opensky-network.org/data','Public scientific datasets and eligibility for historical database access.'),
('9','FAA. Coded Instrument Flight Procedures.','https://www.faa.gov/air_traffic/flight_info/aeronav/digital_products/cifp/','Official contents, ARINC format and 28-day update cycle.'),
('10','Petrovich, M., Black, M. J. and Varol, G. (2023). TMR: Text-to-Motion Retrieval Using Contrastive 3D Human Motion Synthesis. ICCV.','https://arxiv.org/abs/2305.00976','Related cross-modal retrieval methodology for human motion, not aircraft.'),
('11','Wu, Y., Nguyen, H. H., Nguyen, T. and Le, H. (2026). Beyond Tokenization: Direct Timestep Embedding and Contrastive Alignment for Time-Series Question Answering. arXiv preprint.','https://arxiv.org/abs/2606.18986','Recent related alignment work; a preprint rather than evidence of validated aviation transfer.')]
for n,title,url,note in refs:
    a=p(f'[{n}] {title}'); a.paragraph_format.space_after=Pt(2)
    a=p(note); a.paragraph_format.space_after=Pt(2)
    for r in a.runs:r.font.size=Pt(9.5)
    a=doc.add_paragraph(); a.paragraph_format.space_after=Pt(8); link(a,'Source',url)
    for r in a.runs:r.font.size=Pt(9)

page('Repository evidence and reproducibility notes')
p('Local references are relative to C:\\Transformer-rl-retro-propulsion. They document available interfaces and the inspected snapshot; they do not imply that all historical runs share the current code or physics configuration.')
local=[
('L1','simulation/isaac/tvc_env/envs/observations.py','Observation channel order, units, frame conventions and optional wind estimate.'),
('L2','simulation/isaac/apps/run_eval_ppo.py','Trace sampling, stored fields and terminal pre-reset state; see trace_every and construction of the trace dictionary.'),
('L3','simulation/isaac/runs/ppo_diagnostics/','Read-only inventory of the five trajectory.jsonl files listed below. The example landing record is in full28m_seed456/episodes.jsonl.'),
('L4','simulation/isaac/tvc_env/telemetry/episode_export.py','CSV step export plus JSON provenance and episode metadata.'),
('L5','simulation/isaac/apps/run_train_gtrxl.py','Explicitly states that a sequence-aware GTrXL PPO optimizer is absent; environment compatibility smoke only. The available feed-forward model is tvc_env/controllers/ppo_model.py.'),
('L6','simulation/isaac/tvc_env/envs/success_criteria.py','Landing-success helper combines LANDED contact state with a configurable pad-distance threshold.'),
('L7','simulation/isaac/tvc_env/envs/evaluation_contract.py','Physical configuration and source-manifest checks for evaluation comparability.')]
for n,path,note in local:
    a=p(f'[{n}] {path}'); a.paragraph_format.space_after=Pt(2)
    for r in a.runs:r.font.size=Pt(9.5);r.bold=True
    p(note)
table(['Diagnostic run','Episodes','Sampled rows'],[
('full20m_seed123',6,112),('full26m_seed2026_stochastic',6,486),('full28m_seed456',6,524),('stage0_mastered',4,38),('stage1_507904',3,83),('Total inspected',25,1243)],[4.5,1.05,1.35])
p('Recommended manifest fields: source file, run ID, episode ID, checkpoint identifier, seed, task/curriculum stage, configuration hash, source revision, timestamps, sampling interval, schema version, frame/units, action mode, and outcome definition. Missing provenance should remain explicitly missing, not guessed.')
p('Keep query judgments, caption-rule versions and model revisions with the manifest. Training and test partitions should be reproducible from stable IDs. The final artifact should include the corpus manifest, retrieval results and error cases so another researcher can repeat the comparison.')

dest=OUT/'aviation_nlp_embedding_project_ideas.docx'
doc.save(dest)
print(dest)
print('Native equations:', len(doc.element.xpath('//m:oMath')))
